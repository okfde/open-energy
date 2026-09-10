#!/usr/bin/env python3
"""Lädt die Anzahl privater Haushalte je Gemeinde aus dem Zensus 2022
("Regionaltabelle Haushalte", Statistische Ämter des Bundes und der Länder) –
öffentlich und ohne Zugangsdaten abrufbar.

Nur für den deutschlandweiten Balkonkraftwerke-Filter benötigt (Normalisierung
auf Anlagen je 100 Haushalte statt je Einwohner) – nicht für die
Brandenburg-Daten, die weiterhin aus dem Energieportal-WFS kommen.

Der Zensus weist Gemeinden über den 12-stelligen Amtlichen Regionalschlüssel
(ARS) aus, nicht über den 8-stelligen Amtlichen Gemeindeschlüssel (AGS), den
die übrigen Datenquellen dieser Pipeline verwenden. Beide Schlüssel sind
deckungsgleich aufgebaut, der ARS enthält zusätzlich die 4-stellige
Gemeindeverband-Kennziffer: AGS = ARS[0:5] + ARS[9:12] (Land+Regierungsbezirk+
Kreis, dann Gemeinde – der Gemeindeverband-Teil entfällt).

Schreibt scripts/_cache/haushalte_de.json ({ags: haushalte_insgesamt}).
"""
import io
import json
import sys
from pathlib import Path

import openpyxl
import requests

XLSX_URL = "https://www.destatis.de/static/DE/zensus/gitterdaten/Regionaltabelle_Haushalte.xlsx"
SHEET_NAME = "CSV-Haushalte"

CACHE_DIR = Path(__file__).parent / "_cache"
OUT_FILE = CACHE_DIR / "haushalte_de.json"


def ars_to_ags(ars):
    """12-stelliger Amtlicher Regionalschlüssel -> 8-stelliger Amtlicher
    Gemeindeschlüssel (ohne die 4-stellige Gemeindeverband-Kennziffer)."""
    return ars[0:5] + ars[9:12]


def main():
    print(f"Lade Zensus-2022-Haushaltstabelle von {XLSX_URL} …")
    resp = requests.get(XLSX_URL, timeout=180)
    resp.raise_for_status()
    print(f"  {len(resp.content) / 1024 / 1024:.1f} MB heruntergeladen.")

    wb = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"Blatt '{SHEET_NAME}' nicht gefunden – Sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    header = None
    haushalte = {}
    unklar = 0
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = row
            continue
        # Spalten laut Kopfzeile: Berichtszeitpunkt, _RS, Name, Reg_Ebene,
        # 0_Insgesamt_ (Haushalte insgesamt), danach Aufschlüsselungen.
        ars, reg_ebene, insgesamt = row[1], row[3], row[4]
        if reg_ebene != "Gemeinde":
            continue
        if not isinstance(insgesamt, (int, float)):
            # Geheimhaltung (Cell-Key-Methode) markiert einzelne Werte mit
            # "." (unbekannt/geheim) statt einer Zahl – sehr selten auf
            # Gemeindeebene für "Insgesamt", aber robust behandeln.
            unklar += 1
            continue
        haushalte[ars_to_ags(ars)] = int(insgesamt)

    print(f"{len(haushalte)} Gemeinden mit Haushaltszahl verarbeitet.")
    if unklar:
        print(f"Hinweis: {unklar} Gemeinden mit geheimgehaltenem Wert übersprungen.")
    if len(haushalte) < 10000:
        print(
            f"WARNUNG: nur {len(haushalte)} Gemeinden erhalten (erwartet: ~10.800) "
            "– Tabellenstruktur ggf. geändert, bitte prüfen.",
            file=sys.stderr,
        )

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(json.dumps(haushalte, ensure_ascii=False), encoding="utf-8")
    print(f"Geschrieben: {OUT_FILE}")


if __name__ == "__main__":
    main()
